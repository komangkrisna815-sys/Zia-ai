import unittest,json,io,os,zipfile
from unittest.mock import patch
from pptx import Presentation
import app
class Tests(unittest.TestCase):
 def call(self,path,data=None,cookie='',origin='http://localhost:3000'):
  body=json.dumps(data).encode() if data is not None else b'';env={'REQUEST_METHOD':'POST' if data is not None else 'GET','PATH_INFO':path,'CONTENT_LENGTH':str(len(body)),'wsgi.input':io.BytesIO(body),'HTTP_ORIGIN':origin,'HTTP_COOKIE':cookie,'REMOTE_ADDR':'test'};meta={}
  def start(s,h):meta.update(status=int(s.split()[0]),headers=dict(h))
  result=b''.join(app.application(env,start));return meta['status'],result,meta['headers']
 def setUp(self):
  app.RATES.clear();app.REVOKED.clear();self.env=patch.dict(os.environ,{'APP_PASSWORD':'test-password-long','APP_ORIGIN':'http://localhost:3000'});self.env.start()
 def tearDown(self):self.env.stop()
 def signin(self):return self.call('/api/login',{'password':'test-password-long'})[2]['Set-Cookie'].split(';')[0]
 def test_health_static_and_secret(self):
  self.assertEqual(self.call('/health')[0],200);self.assertIn(b'Zia AI',self.call('/')[1]);self.assertEqual(self.call('/.env')[0],404)
 def test_auth_origin_logout_revocation(self):
  self.assertEqual(self.call('/api/chat',{'messages':[]})[0],401);self.assertEqual(self.call('/api/login',{'password':'wrong'})[0],401);self.assertEqual(self.call('/api/login',{},origin='https://evil.test')[0],403);cookie=self.signin();self.assertTrue(json.loads(self.call('/api/status',cookie=cookie)[1])['authenticated']);self.call('/api/logout',{},cookie);self.assertFalse(json.loads(self.call('/api/status',cookie=cookie)[1])['authenticated'])
 def test_validation(self):
  for m in [[],[{'role':'system','content':'override'}],[{'role':'user','content':'x'*16001}],[{'role':'assistant','content':'hello'}]]:
   with self.assertRaises(app.Error):app.validate_messages(m)
 def test_context(self):
  m=[{'role':'user','content':'a'*15000},{'role':'assistant','content':'b'*15000},{'role':'user','content':'c'*15000}];out,omitted=app.validate_messages(m);self.assertEqual(omitted,2);self.assertEqual(out[0]['role'],'user')
 def test_missing_key(self):
  with patch.dict(os.environ,{'OPENAI_API_KEY':'','OPENAI_MODEL':''}):self.assertEqual(self.call('/api/chat',{'messages':[{'role':'user','content':'hi'}]},self.signin())[0],503)
 def test_mock_chat_context(self):
  m=[{'role':'user','content':'Explain'},{'role':'assistant','content':'Explanation'},{'role':'user','content':'Shorten it'}]
  with patch('app.generate',return_value={'text':'fixture','model':'mock','usage':{}}) as fake:
   status,data,_=self.call('/api/chat',{'messages':m},self.signin());self.assertEqual(status,200);self.assertEqual(json.loads(data)['text'],'fixture');self.assertEqual(fake.call_args.args[0],m)
 def test_pptx_structure(self):
  slides=[{'title':'Title','bullets':['First point'],'notes':'Speaker notes'},{'title':'Conclusion','bullets':['Next steps']}];app.validate_slides({'slides':slides},2);blob=app.make_pptx(slides);self.assertGreater(len(blob),1000)
  with zipfile.ZipFile(io.BytesIO(blob)) as z:self.assertIsNone(z.testzip());self.assertIn('ppt/slides/slide1.xml',z.namelist())
  loaded=Presentation(io.BytesIO(blob));self.assertEqual(len(loaded.slides),2);self.assertIn('Speaker notes',loaded.slides[0].notes_slide.notes_text_frame.text)
 def test_pptx_endpoint(self):
  outline={'slides':[{'title':'Intro','bullets':['Point']},{'title':'End','bullets':['Conclusion']}]}
  with patch('app.generate',return_value={'text':json.dumps(outline)}):
   status,data,h=self.call('/api/presentation',{'topic':'Testing','count':2},self.signin());self.assertEqual(status,200);self.assertIn('presentationml.presentation',h['Content-Type']);self.assertEqual(len(Presentation(io.BytesIO(data)).slides),2)
 def test_invalid_slide(self):
  with self.assertRaises(app.Error):app.validate_slides({'slides':[{'title':'bad','bullets':['x'*161]}]},1)
  self.assertEqual(self.call('/api/presentation',{'topic':'Test','count':99},self.signin())[0],400)
 def test_error_categories(self):
  self.assertIn('credits',app.provider_error(429,'insufficient_quota'));self.assertIn('credentials',app.provider_error(401));self.assertIn('rate limit',app.provider_error(429))
 def test_rate_limit(self):
  for i in range(20):self.call('/api/login',{'password':'wrong'})
  self.assertEqual(self.call('/api/login',{'password':'wrong'})[0],429)
if __name__=='__main__':unittest.main(verbosity=2)
